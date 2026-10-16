"""Manual readings ("Ablesungen") at the press of a button.

Each press stores the current scale value with the time of the press. Presses alternate
between start (Masse_0/Zeit_0) and end (Masse_1/Zeit_1) of a measuring step, so one step
becomes one row: Masse_0 | Masse_1 | Zeit_0 | Zeit_1, the same order as columns D:G of the
lab journal, ready to paste there.

Per day there are two files in ``<data folder>/ablesungen/``:

- ``Ablesungen_JJJJ-MM-TT.csv``: append-only log of every press. It is the source of
  truth; the day is rebuilt from it after a restart.
- ``Ablesungen_JJJJ-MM-TT.xlsx``: rewritten from the log after every press. If it is open
  in Excel (locked), the press is still in the CSV and the xlsx catches up on the next
  press after Excel closed it.
"""
from __future__ import annotations

import csv
import os
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional

SLOTS = ("Masse_0", "Masse_1")
CSV_HEADER = ["Nr", "Datum", "Uhrzeit", "Feld", "Wert", "Gueltig", "Rohtext", "ROI", "Run", "Zeit_Messwert"]


@dataclass
class Ablesung:
    nr: int
    when: datetime
    slot: str  # "Masse_0" or "Masse_1"
    value: Optional[float]
    valid: bool
    raw_text: str = ""
    roi: str = ""
    run_id: Optional[int] = None
    reading_time: Optional[datetime] = None


@dataclass
class Step:
    """One row: start and end reading of a measuring step."""
    start: Optional[Ablesung] = None
    end: Optional[Ablesung] = None


class AblesungLog:
    def __init__(self, folder: str | Path, day: Optional[date] = None):
        self.folder = Path(folder)
        self.day = day or date.today()
        self.entries: List[Ablesung] = []
        self._force_new_row = False
        self.load()

    @property
    def csv_path(self) -> Path:
        return self.folder / f"Ablesungen_{self.day.isoformat()}.csv"

    @property
    def xlsx_path(self) -> Path:
        return self.folder / f"Ablesungen_{self.day.isoformat()}.xlsx"

    # -- state -------------------------------------------------------------------------

    @property
    def next_slot(self) -> str:
        if self._force_new_row or not self.entries or self.entries[-1].slot == SLOTS[1]:
            return SLOTS[0]
        return SLOTS[1]

    def start_new_row(self) -> None:
        """The next press starts a new row, even if the last row has no end value yet."""
        self._force_new_row = True

    def steps(self) -> List[Step]:
        steps: List[Step] = []
        for e in self.entries:
            if e.slot == SLOTS[0] or not steps or steps[-1].end is not None:
                steps.append(Step())
            if e.slot == SLOTS[0]:
                steps[-1].start = e
            else:
                steps[-1].end = e
        return steps

    # -- recording ---------------------------------------------------------------------

    def record(
        self,
        value: Optional[float],
        valid: bool = True,
        raw_text: str = "",
        roi: str = "",
        run_id: Optional[int] = None,
        when: Optional[datetime] = None,
        reading_time: Optional[datetime] = None,
    ) -> Ablesung:
        """Stores one press. Raises OSError if the CSV log cannot be written (nothing is
        kept then). A locked xlsx is not an error, see ``write_xlsx``."""
        when = when or datetime.now()
        if when.date() != self.day:
            # Past midnight: continue in the new day's files.
            self.day = when.date()
            self.entries = []
            self.load()
        entry = Ablesung(
            nr=len(self.entries) + 1,
            when=when.replace(microsecond=0),
            slot=self.next_slot,
            value=value,
            valid=valid and value is not None,
            raw_text=raw_text,
            roi=roi,
            run_id=run_id,
            reading_time=reading_time.replace(microsecond=0) if reading_time else None,
        )
        self._append_csv(entry)
        self.entries.append(entry)
        self._force_new_row = False
        return entry

    def _append_csv(self, e: Ablesung) -> None:
        self.folder.mkdir(parents=True, exist_ok=True)
        new = not self.csv_path.exists()
        with open(self.csv_path, "a", newline="", encoding="utf-8-sig" if new else "utf-8") as f:
            w = csv.writer(f, delimiter=";")
            if new:
                w.writerow(CSV_HEADER)
            w.writerow([
                e.nr,
                e.when.date().isoformat(),
                e.when.strftime("%H:%M:%S"),
                e.slot,
                "" if e.value is None else repr(float(e.value)),
                1 if e.valid else 0,
                e.raw_text,
                e.roi,
                "" if e.run_id is None else e.run_id,
                e.reading_time.isoformat(sep=" ") if e.reading_time else "",
            ])
            f.flush()
            os.fsync(f.fileno())

    def load(self) -> None:
        if not self.csv_path.exists():
            return
        with open(self.csv_path, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f, delimiter=";"):
                when = datetime.fromisoformat(f"{row['Datum']} {row['Uhrzeit']}")
                self.entries.append(Ablesung(
                    nr=int(row["Nr"]),
                    when=when,
                    slot=row["Feld"],
                    value=float(row["Wert"]) if row["Wert"] else None,
                    valid=row["Gueltig"] == "1",
                    raw_text=row["Rohtext"],
                    roi=row["ROI"],
                    run_id=int(row["Run"]) if row["Run"] else None,
                    reading_time=datetime.fromisoformat(row["Zeit_Messwert"]) if row["Zeit_Messwert"] else None,
                ))

    # -- Excel -------------------------------------------------------------------------

    def write_xlsx(self) -> bool:
        """Rewrites the day's xlsx from the log. Returns False if the file is locked
        (open in Excel); the data is safe in the CSV and the next call catches up."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Ablesungen"
        ws.append(["Masse_0", "Masse_1", "Zeit_0", "Zeit_1", "", "Run", "Hinweis"])
        warn = PatternFill("solid", fgColor="FFC7CE")
        for step in self.steps():
            s, e = step.start, step.end
            ws.append([
                s.value if s else None,
                e.value if e else None,
                s.when.time() if s else None,
                e.when.time() if e else None,
                None,
                next((x.run_id for x in (s, e) if x and x.run_id is not None), None),
                "; ".join(f"{x.slot} ungültig ({x.raw_text or 'kein Wert'})" for x in (s, e) if x and not x.valid),
            ])
            r = ws.max_row
            ws.cell(r, 3).number_format = ws.cell(r, 4).number_format = "hh:mm:ss"
            for col, x in ((1, s), (2, e)):
                if x and not x.valid:
                    ws.cell(r, col).fill = warn

        log = wb.create_sheet("Protokoll")
        log.append(["Nr", "Zeit", "Feld", "Wert", "Gültig", "Rohtext", "ROI", "Run", "Zeit Messwert"])
        for x in self.entries:
            log.append([x.nr, x.when, x.slot, x.value, "ja" if x.valid else "NEIN", x.raw_text,
                        x.roi, x.run_id, x.reading_time])
            log.cell(log.max_row, 2).number_format = "dd.mm.yyyy hh:mm:ss"
            log.cell(log.max_row, 9).number_format = "dd.mm.yyyy hh:mm:ss"

        for sheet in (ws, log):
            for c in sheet[1]:
                c.font = Font(bold=True)
            sheet.freeze_panes = "A2"
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = 12
        log.column_dimensions["B"].width = log.column_dimensions["I"].width = 20

        self.folder.mkdir(parents=True, exist_ok=True)
        tmp = self.xlsx_path.with_suffix(".tmp.xlsx")
        try:
            wb.save(tmp)
            os.replace(tmp, self.xlsx_path)
        except PermissionError:
            tmp.unlink(missing_ok=True)
            return False
        return True
