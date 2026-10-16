from datetime import date, datetime, time

import pytest
from openpyxl import load_workbook

from instrument_reader.core.ablesung import AblesungLog

DAY = date(2026, 10, 15)


def _at(h, m, s):
    return datetime.combine(DAY, time(h, m, s))


def test_presses_alternate_into_rows_like_the_lab_journal(tmp_path):
    log = AblesungLog(tmp_path, day=DAY)
    assert log.record(9657.45, when=_at(10, 8, 56)).slot == "Masse_0"
    assert log.record(9655.86, when=_at(10, 14, 40)).slot == "Masse_1"
    assert log.record(9655.67, when=_at(10, 15, 0)).slot == "Masse_0"
    assert log.write_xlsx()

    ws = load_workbook(log.xlsx_path)["Ablesungen"]
    assert [c.value for c in ws[1][:4]] == ["Masse_0", "Masse_1", "Zeit_0", "Zeit_1"]
    assert [c.value for c in ws[2][:4]] == [9657.45, 9655.86, time(10, 8, 56), time(10, 14, 40)]
    assert [c.value for c in ws[3][:4]] == [9655.67, None, time(10, 15, 0), None]


def test_day_is_rebuilt_from_csv_after_restart(tmp_path):
    log = AblesungLog(tmp_path, day=DAY)
    log.record(100.5, when=_at(9, 0, 0), run_id=3, roi="Waage")
    log.record(None, valid=False, raw_text="8?.1", when=_at(9, 5, 0))

    again = AblesungLog(tmp_path, day=DAY)
    assert [(e.slot, e.value, e.valid, e.run_id) for e in again.entries] == [
        ("Masse_0", 100.5, True, 3), ("Masse_1", None, False, None)]
    assert again.next_slot == "Masse_0"
    assert again.entries[1].raw_text == "8?.1"


def test_new_row_skips_missing_end(tmp_path):
    log = AblesungLog(tmp_path, day=DAY)
    log.record(1.0, when=_at(9, 0, 0))
    log.start_new_row()
    log.record(2.0, when=_at(9, 1, 0))
    steps = log.steps()
    assert len(steps) == 2 and steps[0].end is None and steps[1].start.value == 2.0


def test_locked_xlsx_keeps_csv_and_catches_up(tmp_path, monkeypatch):
    log = AblesungLog(tmp_path, day=DAY)
    log.record(1.0, when=_at(9, 0, 0))

    def locked(*a, **k):
        raise PermissionError("open in Excel")

    monkeypatch.setattr("instrument_reader.core.ablesung.os.replace", locked)
    assert log.write_xlsx() is False
    assert not list(tmp_path.glob("*.tmp.xlsx"))
    assert log.csv_path.exists()

    monkeypatch.undo()
    log.record(0.5, when=_at(9, 5, 0))
    assert log.write_xlsx()
    assert load_workbook(log.xlsx_path)["Ablesungen"].cell(2, 2).value == 0.5


def test_press_after_midnight_goes_to_new_day(tmp_path):
    log = AblesungLog(tmp_path, day=DAY)
    log.record(1.0, when=_at(23, 59, 59))
    e = log.record(2.0, when=datetime(2026, 10, 16, 0, 0, 5))
    assert e.slot == "Masse_0" and log.csv_path.name == "Ablesungen_2026-10-16.csv"


@pytest.mark.usefixtures("qtbot")
def test_button_stores_current_scale_value(qtbot):
    from instrument_reader.gui.main_window import MainWindow

    w = MainWindow()
    qtbot.addWidget(w)
    w.video_widget.latest_readings = [
        {"roi_name": "Waage", "parsed_value": 9657.8, "is_valid": True, "raw_text": "9657.8", "timestamp": 1_760_000_000.0},
        {"roi_name": "Rotameter", "parsed_value": 3.0, "is_valid": True},
    ]
    w.control_panel.set_scale_roi_names(["Rotameter", "Waage"])
    assert w.control_panel.scale_roi_combo.currentText() == "Waage"

    w.take_ablesung()
    w.take_ablesung()
    steps = w.ablesung_log.steps()
    assert len(steps) == 1 and steps[0].start.value == steps[0].end.value == 9657.8
    assert w.control_panel.ablesung_table.item(0, 0).text() == "9657.8"
    assert w.ablesung_log.xlsx_path.exists()
